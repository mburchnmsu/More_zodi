# List slides of a .pptx (text + pictures, pictures extracted) and extract images from a PDF.
#   poster\.venv\Scripts\python.exe poster\pptx_peek.py OUTDIR deck.pptx template.pdf
import sys, os
from pptx import Presentation
from pptx.util import Emu
import pymupdf

out, deck, pdf = sys.argv[1:4]
os.makedirs(out, exist_ok=True)
p = Presentation(deck)
print("slide size (in):", Emu(p.slide_width).inches, Emu(p.slide_height).inches)
for i, s in enumerate(p.slides, 1):
    print(f"===== slide {i}")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("  [text]", sh.text_frame.text.strip().replace("\n", " | ")[:300])
        if sh.shape_type == 13:  # picture
            fn = f"deck_s{i}_{sh.shape_id}.{sh.image.ext}"
            open(os.path.join(out, fn), "wb").write(sh.image.blob)
            print("  [pic]", fn, sh.image.size)

d = pymupdf.open(pdf)
for pg in d:
    for j, img in enumerate(pg.get_images(full=True)):
        info = d.extract_image(img[0])
        fn = f"tmpl_img{j}.{info['ext']}"
        open(os.path.join(out, fn), "wb").write(info["image"])
        print("[pdf img]", fn, info["width"], info["height"], "bbox", pg.get_image_rects(img[0]))
